import io
import pytest
from openpyxl import load_workbook
from sqlalchemy import insert, select
from src import db

@pytest.mark.parametrize('endpoint',['fleet','models','batteries','forecast','monthly-change','monthly-series','data-quality','sources','summary','filters','fuel-baseline','registrations'])
def test_endpoints_empty(client,endpoint):
    assert client.get('/api/v1/'+endpoint).status_code==200

def test_fleet_filters_pagination(client,populate):
    populate()
    r=client.get('/api/v1/fleet?uf=PR&city=Curitiba&vehicle_year=2026&brand=BYD&model=Dolphin&limit=1').json()
    assert r['total']==1 and r['items'][0]['quantity']==100
    assert r['items'][0]['energy_mwh']==4.49
    assert client.get('/api/v1/fleet?offset=2').json()['items']==[]
    assert client.get('/api/v1/fleet?city=Unknown').json()['total']==0

def test_energy_excludes_unknown(client,populate):
    populate();r=client.get('/api/v1/summary').json()
    assert r['vehicles']==150 and r['energy_mwh']==4.49
    assert r['coverage_pct']==66.7 and r['known_vehicles']==100
    assert r['is_complete_national_coverage'] is False

@pytest.mark.parametrize('scenario,target',[(8,2034),(10,2036),(12,2038)])
def test_forecasts_respect_filters(client,populate,scenario,target):
    populate();r=client.get(f'/api/v1/forecast?uf=PR&scenario={scenario}').json()
    assert r['items'][0]['year']==target and r['items'][0]['vehicles']==100
    assert r['warranty']==[{'year':2034,'vehicles':100}]

def test_unknown_energy_not_zero(client,populate):
    populate();r=client.get('/api/v1/forecast?uf=SP&scenario=8').json()
    assert r['items'][0]['energy_mwh'] is None
    assert r['missing_warranty_vehicles']==50

def test_cascading_locations(client,populate):
    populate();assert client.get('/api/v1/filters?uf=PR').json()['city']==['Curitiba']

@pytest.mark.parametrize('query',['scenario=9','scenario=0','reference=2026-13','vehicle_year=100','capacity_min=-1','capacity_min=50&capacity_max=10','confidence=Z','limit=10001'])
def test_invalid_query(client,query):
    assert client.get('/api/v1/fleet?'+query).status_code==422

def test_previous_month_not_arbitrary(client,populate):
    populate('2026-05');populate('2026-07')
    assert client.get('/api/v1/monthly-change').json()['available'] is False

def test_monthly_delta(client,populate):
    populate('2026-06','UF;MUNICIPIO;MARCA MODELO;ANO FABRICACAO;QUANTIDADE\nPR;Curitiba;I/BYD DOLPHIN GS;2026;80\n');populate()
    r=client.get('/api/v1/monthly-change?uf=PR').json()
    assert r['delta']==20 and r['change_pct']==25 and r['previous_reference']=='2026-06'

@pytest.mark.parametrize('report',['fleet','forecast','monthly-change','data-quality','batteries'])
@pytest.mark.parametrize('fmt',['csv','xlsx'])
def test_exports(client,populate,report,fmt):
    populate();r=client.get(f'/api/v1/export/{report}?format={fmt}&uf=PR')
    assert r.status_code==200 and 'attachment' in r.headers['content-disposition']
    if fmt=='xlsx':assert load_workbook(io.BytesIO(r.content)).active.max_row>=1
    else:assert r.content.startswith(b'\xef\xbb\xbf')

def test_admin_protected(client,monkeypatch):
    monkeypatch.delenv('ADMIN_API_KEY',raising=False)
    assert client.post('/api/v1/admin/etl/run',json={}).status_code==503
    monkeypatch.setenv('ADMIN_API_KEY','test-secret')
    assert client.post('/api/v1/admin/etl/run',json={}).status_code==401
    monkeypatch.setenv('ALLOW_SOURCE_DOWNLOAD','false')
    assert client.post('/api/v1/admin/etl/run',json={},headers={'X-Admin-Key':'test-secret'}).status_code==409

def test_manual_match_cannot_change_history(client,populate,monkeypatch):
    populate();monkeypatch.setenv('ADMIN_API_KEY','test')
    r=client.post('/api/v1/admin/matches',headers={'X-Admin-Key':'test'},json={'raw_name':'BYD DOLPHIN GS EXTRA','vehicle_year':2026,'model_id':'byd-dolphin-gs-2026','approver':'Reviewer','notes':'Verificado na ficha brasileira da versão.'})
    assert r.status_code==200
    assert client.get('/api/v1/summary').json()['vehicles']==150

def test_old_year_match_rejected(client,monkeypatch):
    monkeypatch.setenv('ADMIN_API_KEY','test')
    r=client.post('/api/v1/admin/matches',headers={'X-Admin-Key':'test'},json={'raw_name':'BYD DOLPHIN GS','vehicle_year':2023,'model_id':'byd-dolphin-gs-2026','approver':'Reviewer','notes':'Tentativa de ano fora da vigência.'})
    assert r.status_code==422

def test_new_snapshot_invalidates_default_reference_cache(client,populate):
    populate('2026-06')
    assert client.get('/api/v1/summary').json()['vehicles']==150
    populate('2026-07',text='UF;MUNICIPIO;MARCA MODELO;ANO FABRICACAO;QUANTIDADE\nPR;Curitiba;I/BYD DOLPHIN GS;2026;101\n')
    response=client.get('/api/v1/summary').json()
    assert response['reference']=='2026-07' and response['vehicles']==101
    assert client.get('/api/v1/summary?reference=2026-06').json()['vehicles']==150

def test_propulsion_filter_and_official_coverage(client,populate):
    populate()
    summary=client.get('/api/v1/summary').json()
    assert summary['by_propulsion'][0]['name']=='BEV'
    assert summary['official_coverage']['items']
    assert client.get('/api/v1/fleet?propulsion=BEV').json()['total']==2
    assert client.get('/api/v1/fleet?propulsion=PHEV').json()['total']==0
    quality=client.get('/api/v1/data-quality').json()
    assert quality['unclassified'][0]['raw_name']=='TOYOTA COROLLA'
    assert quality['unclassified_other_vehicles']==999
    series=client.get('/api/v1/monthly-series').json()
    assert series['items'][0]['vehicles']==150
    assert any(r['provider']=='Fenabrave' for r in series['registrations'])
    assert summary['official_coverage']['official_total']
    assert summary['official_ruler_pct'] is not None
    assert summary['energy_by_propulsion']

def test_orphan_runs_are_failed_on_startup(database):
    from src.api.main import fail_orphan_runs
    from src.connectors.public import now
    with database.begin() as c:
        c.execute(insert(db.runs).values(id='orphan',reference='2026-07',started_at=now(),status='EXECUTANDO',stage='Baixando',details={}))
    fail_orphan_runs()
    with database.connect() as c:
        assert c.execute(select(db.runs.c.status).where(db.runs.c.id=='orphan')).scalar()=='ERRO'
import pytest

@pytest.fixture(autouse=True)
def legacy_contract_mode(monkeypatch):
    """Exercise legacy implementations before cutover, independently of write shutdown."""
    monkeypatch.setenv('I9_LEGACY_WRITES','true')
