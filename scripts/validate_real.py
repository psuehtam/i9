"""Reproducible HTTP checks against published SENATRAN snapshots."""
import csv,io,json,os,time
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
import httpx

ROOT=Path(__file__).resolve().parents[1]
REFERENCE=os.getenv('I9_VALIDATE_REFERENCE','2026-07')

def main():
    checks=[]
    with httpx.Client(base_url='http://127.0.0.1:8000',timeout=60) as c:
        def req(path):
            path+=('&' if '?' in path else '?')+'reference='+REFERENCE
            started=time.perf_counter();r=c.get('/api/v1/'+path);r.raise_for_status()
            checks.append({'endpoint':path,'http':r.status_code,'seconds':round(time.perf_counter()-started,3),'response_bytes':len(r.content)})
            return r
        s=req('summary').json();assert s['vehicles']>0
        if REFERENCE=='2026-07':
            assert s['vehicles']==222373
        assert s['official_coverage']['official_total']
        assert s.get('energy_by_propulsion') is not None
        m=req('monthly-change').json()
        if REFERENCE=='2026-07':
            assert m['previous']==205816 and m['delta']==16557
        for scenario in (8,10,12):
            forecast=req('forecast?scenario='+str(scenario)).json()
            assert sum(x['vehicles'] for x in forecast['items'])+forecast['missing_year_vehicles']==s['vehicles']
        filtered=req('fleet?uf=PR&city=Curitiba&vehicle_year=2026&limit=1000').json()
        assert filtered['items'] and all(r['uf']=='PR' and r['city']=='Curitiba' and r['vehicle_year']==2026 for r in filtered['items'])
        req('filters?uf=PR');quality=req('data-quality').json()
        for endpoint in ('models','batteries','sources'):req(endpoint)
        r=req('export/fleet?format=csv')
        rows=list(csv.DictReader(io.StringIO(r.content.decode('utf-8-sig')),delimiter=';'))
        assert sum(int(x['quantity']) for x in rows)==s['vehicles']
        excel=req('export/fleet?format=xlsx')
        from openpyxl import load_workbook
        wb=load_workbook(io.BytesIO(excel.content),read_only=True)
        sheet=wb.active;iterator=sheet.iter_rows(values_only=True);headers=list(next(iterator));position=headers.index('quantity')
        assert sum(r[position] for r in iterator)==s['vehicles'];wb.close()
        print('Requisições com dados reais e exportações conferidas.',flush=True)
        started=time.perf_counter()
        with ThreadPoolExecutor(max_workers=8) as pool:
            statuses=list(pool.map(lambda p:c.get('/api/v1/'+p+'?reference='+REFERENCE).status_code,['summary','forecast','data-quality','filters']*4))
        assert all(code==200 for code in statuses)
        concurrency_seconds=round(time.perf_counter()-started,3)
        settings=dict(line.split('=',1) for line in (ROOT/'.env').read_text(encoding='utf-8').splitlines() if line and not line.startswith('#'))
        started=time.perf_counter();r=c.post('/api/v1/admin/etl/run',json={},headers={'X-Admin-Key':settings['ADMIN_API_KEY']});r.raise_for_status();run_id=r.json()['run_id']
        for _ in range(100):
            run=c.get('/api/v1/etl/runs/'+run_id).json()
            if run['status']!='EXECUTANDO':break
            time.sleep(.2)
        assert run['status'] in ('SEM_NOVIDADE','SUCESSO')
        result={'requests':checks,'concurrent_requests':len(statuses),'concurrent_ok':True,'concurrent_batch_seconds':concurrency_seconds,
                'etl_status':run['status'],'etl_reference':run.get('reference'),'no_new_month_seconds':round(time.perf_counter()-started,3),'summary':s,'monthly':{k:v for k,v in m.items() if k!='items'},
                'missing_year_vehicles':quality['missing_year_vehicles'],'missing_location_vehicles':quality['missing_location_vehicles'],
                'export_csv_vehicles':sum(int(x['quantity']) for x in rows),'export_xlsx_verified':True}
        (ROOT/'docs/real-data-validation.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
        print({k:v for k,v in result.items() if k not in ('summary','requests')},flush=True)

if __name__=='__main__':main()
